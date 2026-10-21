"""Exposure handoff workbook: layout, status colouring, editable columns."""
import io
import os
from datetime import datetime
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from exposure import rules, signatures, workbook

SIGNATURE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                              "signatures", "php-rce-2026-10.yaml")


@pytest.fixture
def signature():
    return signatures.load_signature_file(SIGNATURE_PATH)


def host(name, verdict, **kw):
    base = dict(host=name, zone_name="example.com", verdict=verdict, detected_version=None, version_status="unknown",
                hosting_provider=None, origin_ip=None, owner=None, contact=None, decision=None, notes=None,
                verify_error=None, summary={"matching_requests": 10, "origin_2xx": 5, "distinct_ips_2xx": 3})
    base.update(kw)
    return SimpleNamespace(**base)


def finding(name, rule_id, verdict, trusted=True):
    return SimpleNamespace(host=name, rule_id=rule_id, verdict=verdict, confidence="inferred",
                           origin_status_trusted=trusted, detail=f"{rule_id} fired on {name}", evidence={"n": 1})


def build(signature, **run_kw):
    run = SimpleNamespace(id=7, status="success", signature_name="php-rce-2026-10", signature_version="2",
                          window_start=datetime(2026, 10, 14), window_end=datetime(2026, 10, 21),
                          verify_enabled=False, coverage_note="sampled", edge_coverage={"example.com": "disabled"},
                          stats={"zones": 2, "hosts": 3, "graphql_queries": 6,
                                 "verdicts": {"likely_affected": 1, "misrouted": 1, "catch_all": 1}})
    for k, v in run_kw.items():
        setattr(run, k, v)
    hosts = [
        host("a.example.com", rules.VERDICT_LIKELY, version_status="eol", detected_version="7.4.3", decision="patch", owner="web"),
        host("b.example.com", rules.VERDICT_MISROUTED),
        host("c.example.com", rules.VERDICT_CATCH_ALL),
    ]
    findings = [finding("a.example.com", "EXP-001", rules.VERDICT_LIKELY),
                finding("b.example.com", "EXP-003", rules.VERDICT_MISROUTED, trusted=False)]
    wb = workbook.build_exposure_workbook(run, hosts, findings, [], signature, now=datetime(2026, 10, 21, 3, 0))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf)


def test_sheet_order_and_cover_content(signature):
    wb = build(signature)
    assert wb.sheetnames == ["Summary", "Action Items", "Evidence", "Cleared", "Intel Gaps", "Mitigations"]
    ws = wb["Summary"]
    # Logo image anchored over the top rows (rendered by branding.wordmark_png)
    assert len(ws._images) == 1
    assert ws._images[0].anchor._from.row == workbook.LOGO_ROWS[0] - 1
    assert ws["B4"].value == "Exposure analysis handoff"
    assert "run #7" in ws["B5"].value
    # KPI tiles: label / value / sub, headline in brand orange
    assert ws["B7"].value == "LIKELY AFFECTED" and ws["B8"].value == 1
    assert ws["B8"].font.color.rgb.endswith(workbook.BRAND_ORANGE)
    assert ws["C8"].font.color.rgb.endswith(workbook.TEXT)
    assert ws["B9"].value == "all decided"            # the one likely-affected host has decision=patch
    assert ws["C9"].value == "1 without a decision"   # the misrouted host does not
    # Managed WAF coverage is summarised by state, disabled in the medium tint
    labels = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3) for r in range(1, ws.max_row + 1)}
    assert labels["disabled (1 zone)"].value == "example.com"
    assert labels["disabled (1 zone)"].fill.start_color.rgb.endswith(workbook.TIERS["medium"][0])
    assert "enabled (0 zones)" not in labels
    assert not ws.sheet_view.showGridLines


def test_action_items_are_filterable_editable_and_colour_coded(signature):
    ws = build(signature)["Action Items"]
    headers = [c.value for c in ws[1]]
    assert headers[0] == "Host" and headers[11] == "Decision"
    assert ws.auto_filter.ref == "A1:N3"
    assert ws.freeze_panes == "B2"
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {"a.example.com", "b.example.com"}  # catch-all is not an action item
    # verdict + version status cells carry their tier colours
    a = ws[2] if ws["A2"].value == "a.example.com" else ws[3]
    assert a[2].fill.start_color.rgb.endswith(workbook.TIERS["critical"][0])
    assert a[6].value == "eol" and a[6].fill.start_color.rgb.endswith(workbook.TIERS["critical"][0])
    # editable columns are tinted and the decision column has a dropdown
    assert a[9].fill.start_color.rgb.endswith(workbook.TIERS["info"][0])
    dvs = list(ws.data_validations.dataValidation)
    assert len(dvs) == 1 and dvs[0].type == "list" and "patch" in dvs[0].formula1
    assert "L2:L3" in str(dvs[0].sqref)
    # undecided rows get a conditional-format highlight
    cf_ranges = [str(r) for r in ws.conditional_formatting]
    assert any("L2:L3" in r for r in cf_ranges)


def test_cleared_numbers_and_gap_tiers(signature):
    wb = build(signature)
    cleared = wb["Cleared"]
    row = next(r for r in cleared.iter_rows(min_row=2) if r[0].value == "c.example.com")
    assert row[3].value == 10 and row[3].number_format == "#,##0" and row[3].alignment.horizontal == "right"
    assert row[2].fill.start_color.rgb.endswith(workbook.TIERS["low"][0])

    gaps = wb["Intel Gaps"]
    kinds = {r[1].value: r[1] for r in gaps.iter_rows(min_row=2)}
    assert "Passive only" in kinds and "Owner unknown" in kinds and "Untrusted origin status" not in kinds
    assert kinds["Passive only"].fill.start_color.rgb.endswith(workbook.TIERS["low"][0])
    assert kinds["Owner unknown"].fill.start_color.rgb.endswith(workbook.TIERS["medium"][0])


def test_mitigations_sheet_lists_every_expression_in_monospace(signature):
    ws = build(signature)["Mitigations"]
    cells = [c for row in ws.iter_rows(min_col=2, max_col=2) for c in row if c.value]
    mono = [c for c in cells if c.font.name == workbook.MONO]
    exprs = [c for c in mono if "http.request" in str(c.value)]
    expected = sum(len(m.expressions) for m in signature.mitigations)
    assert len(exprs) == expected
    assert all(c.alignment.wrap_text for c in mono)
    # plan notes sit next to expressions but are prose, not code
    notes = [c for c in cells if c.font.name != workbook.MONO and "http.request" in str(c.value)]
    assert notes and all(not c.font.bold for c in notes)
    titles = [c.value for c in cells if c.font.bold and c.font.size == workbook.SIZE_TITLE]
    assert len(titles) == len(signature.mitigations)
    assert any("session.upload_progress.enabled = 0" in str(c.value) for c in cells)


def test_type_scale_has_a_12pt_floor_and_one_body_face(signature):
    wb = build(signature)
    sizes, faces = set(), set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or c.font is None:
                    continue
                sizes.add(c.font.size or 11)
                faces.add(c.font.name)
    assert min(sizes) >= workbook.SIZE_BODY, sorted(sizes)
    # Only the workbook default face and the monospace face for code / JSON
    assert faces <= {None, "Calibri", workbook.MONO}, faces
    # Hostnames are body text, not small mono: the complaint that drove the scale
    host_cell = wb["Action Items"]["A2"]
    assert host_cell.font.name != workbook.MONO and host_cell.font.size == workbook.SIZE_BODY
    assert wb["Evidence"]["G2"].font.name == workbook.MONO and wb["Evidence"]["G2"].font.size == workbook.SIZE_BODY


def test_formula_injection_is_defused_everywhere(signature):
    wb = build(signature, coverage_note="=HYPERLINK(\"http://evil\")")
    ws = wb["Summary"]
    values = [ws.cell(row=r, column=3).value for r in range(1, ws.max_row + 1)]
    assert any(v == "'=HYPERLINK(\"http://evil\")" for v in values)


def test_cover_falls_back_to_rich_text_without_an_image_backend(signature, monkeypatch):
    import branding
    monkeypatch.setattr(branding, "wordmark_png", lambda **kw: None)
    wb = build(signature)
    ws = wb["Summary"]
    assert ws._images == []
    first, last = workbook.LOGO_ROWS
    assert ws.cell(row=first, column=2).font.size == workbook.SIZE_WORDMARK
    assert ws.cell(row=last, column=2).value == branding.TAGLINE.upper()

"""Exports (CSV/XLSX/PDF) and report presets."""
import io

import main


def seed_minimal(db):
    db.add(main.Zone(zone_id="z1", name="a.com", status="active"))
    db.add(main.DNSRecord(record_id="r1", zone_id="z1", name="www.a.com", type="A",
                          content="192.0.2.1", proxied=True, ttl=1))
    db.commit()


def test_export_csv_default(client, db, admin_user, auth_headers):
    seed_minimal(db)
    resp = client.get("/api/export/dns-records", headers=auth_headers)
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    assert ".csv" in resp.headers["content-disposition"]
    body = resp.text
    assert "www.a.com" in body and "192.0.2.1" in body


def test_export_xlsx(client, db, admin_user, auth_headers):
    from openpyxl import load_workbook
    seed_minimal(db)
    resp = client.get("/api/export/dns-records?format=xlsx", headers=auth_headers)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][3] == "Name"  # header row
    assert rows[1][3] == "www.a.com"


def test_export_rejects_bad_format(client, db, admin_user, auth_headers):
    resp = client.get("/api/export/zones?format=exe", headers=auth_headers)
    assert resp.status_code == 422


def test_formula_injection_sanitized(client, db, admin_user, auth_headers):
    db.add(main.Zone(zone_id="z1", name="a.com", status="active"))
    db.add(main.DNSRecord(record_id="r1", zone_id="z1", name="www.a.com", type="TXT",
                          content="=HYPERLINK(\"http://evil\")", ttl=1))
    db.commit()
    resp = client.get("/api/export/dns-records", headers=auth_headers)
    assert "'=HYPERLINK" in resp.text  # leading quote defuses the formula


def test_summary_pdf(client, db, admin_user, auth_headers):
    seed_minimal(db)
    resp = client.get("/api/export/summary", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content[:5] == b"%PDF-"


def test_preset_crud_and_audit(client, db, admin_user, auth_headers):
    resp = client.post("/api/report-presets", headers=auth_headers, json={
        "name": "Weekly traffic", "report_id": "traffic", "format": "xlsx", "params": {"hours": 168}
    })
    assert resp.status_code == 200
    preset_id = resp.json()["id"]

    listed = client.get("/api/report-presets", headers=auth_headers).json()
    assert [p["name"] for p in listed] == ["Weekly traffic"]
    assert listed[0]["params"] == {"hours": 168}

    resp = client.delete(f"/api/report-presets/{preset_id}", headers=auth_headers)
    assert resp.status_code == 200
    assert client.get("/api/report-presets", headers=auth_headers).json() == []

    actions = [a.action for a in db.query(main.AuditLog).all()]
    assert "PRESET_CREATED" in actions and "PRESET_DELETED" in actions


def test_preset_rejects_unknown_report(client, admin_user, auth_headers):
    resp = client.post("/api/report-presets", headers=auth_headers, json={
        "name": "Bad", "report_id": "not-a-report", "format": "csv"
    })
    assert resp.status_code == 422

"""Exposure Analysis end-to-end: background run with canned Cloudflare
responses, API endpoints, workbook export and diff.

Acceptance criteria 8 (sampling flagged) and 9 (handoff persistence) from
docs/exposure-analysis/SPEC.md.
"""
import io
from datetime import datetime, timedelta

import pytest

import main
from exposure import collector, enrich, verify
from exposure.collector import DATASET


@pytest.fixture()
def config(db, admin_user):
    cfg = main.CloudflareConfig(name="Prod", api_token=main.encrypt_token("cf-token"), account_id="acct",
                                is_active=True, created_by=admin_user.id)
    db.add(cfg)
    db.commit()
    db.refresh(cfg)
    db.add(main.Zone(zone_id="z1", config_id=cfg.id, name="example.com", status="active"))
    db.add(main.DNSRecord(record_id="r1", zone_id="z1", name="realwp.example.com", type="A", content="203.0.113.37", proxied=True))
    db.add(main.DNSRecord(record_id="r2", zone_id="z1", name="webmail.example.com", type="A", content="104.21.15.127", proxied=True))
    db.add(main.DNSRecord(record_id="r3", zone_id="z1", name="webmail.example.com", type="A", content="198.51.100.44", proxied=True))
    db.commit()
    return cfg


def group(host, path, method="GET", edge=200, origin=200, count=1, ip=None):
    dims = {"clientRequestHTTPHost": host, "clientRequestPath": path, "clientRequestHTTPMethodName": method,
            "edgeResponseStatus": edge, "originResponseStatus": origin}
    if ip:
        dims["clientIP"] = ip
    return {"count": count, "avg": {"sampleInterval": 10}, "dimensions": dims}


@pytest.fixture()
def offline(monkeypatch):
    """Replace every outbound call the run makes with canned data"""
    paths = [
        group("realwp.example.com", "/wp-login.php", count=193),
        group("spa.example.com", "/phpinfo.php~"),
        group("spa.example.com", "/cakephp/.env"),
        group("spa.example.com", "/index.php", count=60),
        group("webmail.example.com", "/cakephp/.env", count=2),
    ]
    ips = [{"count": 1, "dimensions": {"clientRequestHTTPHost": "realwp.example.com", "clientIP": f"203.0.113.{i}"}}
           for i in range(1, 25)]
    ips += [{"count": 1, "dimensions": {"clientRequestHTTPHost": "spa.example.com", "clientIP": "198.51.100.9"}}]
    calls = [{"count": 37, "dimensions": {"clientRequestHTTPHost": "realwp.example.com", "clientRequestPath": "/wp-cron.php",
                                          "clientRequestHTTPMethodName": "POST", "clientIP": "203.0.113.37"}}]

    async def fake_graphql(self, query, variables=None):
        def wrap(groups):
            return {"data": {"viewer": {"zones": [{DATASET: groups}]}}}
        if "settings" in query:
            return {"data": {"viewer": {"zones": [{"settings": {DATASET: {"maxDuration": 86400 * 7}}}]}}}
        if "originResponseStatus_geq" in query:
            return wrap(ips)
        if "OR: [" not in query and "clientRequestPath_in" in query:
            return wrap(calls)
        return wrap(paths)

    async def fake_ranges(client=None, force=False):
        return ["104.21.0.0/16", "172.67.0.0/16"], None

    async def fake_hosting(ips_, client=None, concurrency=4):
        return {ip: {"provider": "GoDaddy.com, LLC", "name": "GO-DADDY-COM-LLC"} for ip in ips_}, None

    async def fake_feed(url, client=None):
        return {"8": {"8.3": {"version": "8.3.35"}, "8.4": {"version": "8.4.18"}}}, None

    async def fake_waf(self, zone_id):
        return "disabled"

    monkeypatch.setattr(main.CloudflareClient, "_graphql_request", fake_graphql)
    monkeypatch.setattr(main.CloudflareClient, "get_managed_waf_state", fake_waf)
    monkeypatch.setattr(enrich, "get_cloudflare_ranges", fake_ranges)
    monkeypatch.setattr(enrich, "lookup_hosting", fake_hosting)
    monkeypatch.setattr(enrich, "fetch_version_feed", fake_feed)
    monkeypatch.setattr(main, "EXPOSURE_VERIFY_DELAY_SECONDS", 0)
    main._update_exposure_progress(active=False, result=None, run_id=None)


def start_run(db, config, verify_enabled=False, window_days=7):
    end = datetime.utcnow().replace(second=0, microsecond=0)
    run = main.ExposureRun(config_id=config.id, signature_name="php-rce-2026-10", signature_version="1",
                           window_start=end - timedelta(days=window_days), window_end=end,
                           status="running", verify_enabled=verify_enabled)
    db.add(run)
    db.commit()
    db.refresh(run)
    return run


def test_passive_run_classifies_and_flags_sampling(db, config, offline):
    run = start_run(db, config)
    main.run_exposure_in_background(run.id, user_id=None, ip_address=None)
    db.expire_all()
    run = db.get(main.ExposureRun, run.id)

    assert run.status == "success", run.error
    # 8. Sampling flagged
    assert run.coverage_note and "estimates" in run.coverage_note
    assert run.edge_coverage == {"example.com": "disabled"}
    assert run.stats["verdicts"] == {"likely_affected": 1, "catch_all": 1, "misrouted": 1}

    hosts = {h.host: h for h in db.query(main.ExposureHost).filter(main.ExposureHost.run_id == run.id).all()}
    assert hosts["realwp.example.com"].verdict == "likely_affected"
    assert hosts["realwp.example.com"].hosting_provider == "GoDaddy.com, LLC"
    assert hosts["realwp.example.com"].origin_ip == "203.0.113.37"
    assert hosts["realwp.example.com"].zone_name == "example.com"
    assert hosts["realwp.example.com"].summary["distinct_ips_2xx"] == 24
    assert hosts["spa.example.com"].verdict == "catch_all"
    assert hosts["webmail.example.com"].verdict == "misrouted"
    # Zone-level coverage finding exists but is not a host row
    assert "example.com" not in hosts
    rules_seen = {(f.host, f.rule_id) for f in db.query(main.ExposureFinding).filter(main.ExposureFinding.run_id == run.id)}
    assert ("example.com", "EXP-006") in rules_seen
    assert ("webmail.example.com", "EXP-003") in rules_seen
    assert not db.query(main.ExposureProbe).filter(main.ExposureProbe.run_id == run.id).count()

    audit = db.query(main.AuditLog).filter(main.AuditLog.action == "EXPOSURE_RUN_SUCCESS").first()
    assert audit and audit.resource_type == "exposure"
    assert main._get_exposure_progress()["active"] is False


def test_handoff_fields_survive_next_run(db, config, offline):
    first = start_run(db, config)
    main.run_exposure_in_background(first.id, None, None)
    row = db.query(main.ExposureHost).filter(main.ExposureHost.run_id == first.id,
                                             main.ExposureHost.host == "realwp.example.com").first()
    row.owner, row.contact, row.decision, row.notes = "web team", "web@example.com", "patch", "cPanel host"
    db.commit()

    second = start_run(db, config)
    main.run_exposure_in_background(second.id, None, None)
    db.expire_all()
    carried = db.query(main.ExposureHost).filter(main.ExposureHost.run_id == second.id,
                                                 main.ExposureHost.host == "realwp.example.com").first()
    # 9. Handoff persistence
    assert (carried.owner, carried.contact, carried.decision, carried.notes) == \
        ("web team", "web@example.com", "patch", "cPanel host")
    untouched = db.query(main.ExposureHost).filter(main.ExposureHost.run_id == second.id,
                                                   main.ExposureHost.host == "spa.example.com").first()
    assert untouched.owner is None


def test_verify_run_records_probes_and_upgrades_confidence(db, config, offline, monkeypatch):
    probed = []

    async def fake_verify(host, signature, ssrf_check, **kwargs):
        probed.append(host)
        assert ssrf_check is main._ssrf_check
        if host == "realwp.example.com":
            return {"host": host, "resolver": "cloudflare-doh(1.1.1.1)", "resolved_ip": "8.8.8.8", "resolve_error": None,
                    "blocked": False, "version_header": "PHP/8.3.33", "server_header": "Apache",
                    "probes": [{"kind": "baseline", "path": "/abc", "status": 404, "size": 10, "content_type": "text/html",
                                "title": None, "sha256_prefix": "a" * 16, "secret_match_count": 0, "debug_marker_count": 0,
                                "error": None, "elapsed_ms": 5, "location": None},
                               {"kind": "suspect", "path": "/config.php.bak", "status": 200, "size": 300, "content_type": "text/plain",
                                "title": None, "sha256_prefix": "b" * 16, "secret_match_count": 4, "debug_marker_count": 0,
                                "error": None, "elapsed_ms": 7, "location": None}]}
        return {"host": host, "resolver": None, "resolved_ip": None, "resolve_error": "public resolver returned no A/AAAA answers; system resolver also failed",
                "blocked": False, "version_header": None, "server_header": None, "probes": []}

    monkeypatch.setattr(verify, "verify_host", fake_verify)
    run = start_run(db, config, verify_enabled=True)
    main.run_exposure_in_background(run.id, None, None)
    db.expire_all()
    run = db.get(main.ExposureRun, run.id)
    assert run.status == "success", run.error
    assert set(probed) == {"realwp.example.com", "spa.example.com", "webmail.example.com"}

    host = db.query(main.ExposureHost).filter(main.ExposureHost.run_id == run.id,
                                              main.ExposureHost.host == "realwp.example.com").first()
    assert (host.detected_version, host.version_status) == ("8.3.33", "behind")
    assert host.resolver == "cloudflare-doh(1.1.1.1)" and host.resolved_ip == "8.8.8.8"
    findings = {f.rule_id: f for f in db.query(main.ExposureFinding).filter(
        main.ExposureFinding.run_id == run.id, main.ExposureFinding.host == "realwp.example.com")}
    assert findings["EXP-001"].confidence == "verified"
    assert findings["EXP-005"].evidence["latest"] == "8.3.35"
    assert findings["EXP-008"].evidence["secret_match_count"] == 4
    probes = db.query(main.ExposureProbe).filter(main.ExposureProbe.run_id == run.id).all()
    assert {p.path for p in probes} == {"/abc", "/config.php.bak"}
    unresolved = db.query(main.ExposureHost).filter(main.ExposureHost.run_id == run.id,
                                                    main.ExposureHost.host == "spa.example.com").first()
    assert "resolver" in unresolved.verify_error
    assert "Active verify probed" in run.coverage_note


def test_failed_run_is_recorded(db, config, offline, monkeypatch):
    async def boom(self, query, variables=None):
        raise RuntimeError("graphql down")
    monkeypatch.setattr(main.CloudflareClient, "_graphql_request", boom)
    run = start_run(db, config)
    main.run_exposure_in_background(run.id, None, None)
    db.expire_all()
    run = db.get(main.ExposureRun, run.id)
    assert run.status == "failed" and "Every zone failed" in run.error
    assert main._get_exposure_progress()["result"] == "failed"
    assert db.query(main.AuditLog).filter(main.AuditLog.action == "EXPOSURE_RUN_FAILED").count() == 1


# ---- API surface

def test_signatures_endpoint(client, auth_headers):
    resp = client.get("/api/exposure/signatures", headers=auth_headers)
    assert resp.status_code == 200
    sig = resp.json()[0]
    assert sig["name"] == "php-rce-2026-10"
    assert sig["min_distinct_ips"] == 10
    # Secret patterns are not exposed to the UI, only their count
    assert "secret_patterns" not in sig


def test_start_run_requires_admin_and_rejects_parallel(client, db, config, auth_headers, offline, monkeypatch):
    # Keep the background task from actually running so the 409 path is observable
    monkeypatch.setattr(main, "run_exposure_in_background", lambda *a, **k: None)

    user = main.User(email="u@test.local", username="user", hashed_password=main.get_password_hash("UserPass123!x"),
                     is_admin=False, is_active=True)
    db.add(user)
    db.commit()
    user_headers = {"Authorization": f"Bearer {main.create_access_token({'sub': 'user'})}"}
    body = {"config_id": config.id, "signature_name": "php-rce-2026-10", "window_days": 7, "verify": False}
    assert client.post("/api/exposure/runs", json=body, headers=user_headers).status_code == 403

    resp = client.post("/api/exposure/runs", json=body, headers=auth_headers)
    assert resp.status_code == 202, resp.text
    run = resp.json()
    assert run["status"] == "running" and run["signature_name"] == "php-rce-2026-10"
    assert db.query(main.AuditLog).filter(main.AuditLog.action == "EXPOSURE_RUN_STARTED").count() == 1

    assert client.post("/api/exposure/runs", json=body, headers=auth_headers).status_code == 409
    main._update_exposure_progress(active=False)

    bad = dict(body, signature_name="nope")
    assert client.post("/api/exposure/runs", json=bad, headers=auth_headers).status_code == 404
    assert client.get("/api/exposure/runs/progress", headers=auth_headers).json()["run_id"] == run["id"]


def test_run_detail_patch_export_and_diff(client, db, config, auth_headers, offline):
    first = start_run(db, config)
    main.run_exposure_in_background(first.id, None, None)
    second = start_run(db, config)
    main.run_exposure_in_background(second.id, None, None)

    runs = client.get("/api/exposure/runs", headers=auth_headers).json()
    assert [r["id"] for r in runs] == [second.id, first.id]

    detail = client.get(f"/api/exposure/runs/{second.id}", headers=auth_headers).json()
    assert detail["run"]["status"] == "success"
    assert detail["signature"]["title"].startswith("PHP critical RCE")
    assert {h["host"] for h in detail["hosts"]} == {"realwp.example.com", "spa.example.com", "webmail.example.com"}
    assert any(f["rule_id"] == "EXP-001" and f["verdict"] == "likely_affected" for f in detail["findings"])

    resp = client.patch(f"/api/exposure/hosts/{second.id}/realwp.example.com",
                        json={"owner": "web team", "decision": "patch"}, headers=auth_headers)
    assert resp.status_code == 200 and resp.json()["owner"] == "web team"
    assert client.patch(f"/api/exposure/hosts/{second.id}/realwp.example.com",
                        json={"decision": "yolo"}, headers=auth_headers).status_code == 422
    assert client.patch(f"/api/exposure/hosts/{second.id}/missing.example.com",
                        json={"owner": "x"}, headers=auth_headers).status_code == 404
    audit = db.query(main.AuditLog).filter(main.AuditLog.action == "EXPOSURE_HOST_UPDATED").first()
    assert audit.details["changes"] == {"owner": "web team", "decision": "patch"}

    resp = client.get(f"/api/exposure/runs/{second.id}/export.xlsx", headers=auth_headers)
    assert resp.status_code == 200
    assert "exposure_php-rce-2026-10" in resp.headers["content-disposition"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Action Items", "Evidence", "Cleared", "Intel Gaps", "Mitigations"]
    action_hosts = [row[0] for row in wb["Action Items"].iter_rows(min_row=2, values_only=True)]
    assert set(action_hosts) == {"realwp.example.com", "webmail.example.com"}
    cleared_hosts = [row[0] for row in wb["Cleared"].iter_rows(min_row=2, values_only=True)]
    assert cleared_hosts == ["spa.example.com"]
    gaps = [row[1] for row in wb["Intel Gaps"].iter_rows(min_row=2, values_only=True)]
    assert "Passive only" in gaps and "Version unknown" in gaps

    diff = client.get(f"/api/exposure/runs/{second.id}/diff/{first.id}", headers=auth_headers).json()
    assert diff["new_hosts"] == [] and diff["verdict_changes"] == []
    assert {d["host"] for d in diff["decisions_open"]} == {"webmail.example.com"}  # realwp decided above

    assert client.delete(f"/api/exposure/runs/{first.id}", headers=auth_headers).status_code == 200
    assert client.get(f"/api/exposure/runs/{first.id}", headers=auth_headers).status_code == 404
    assert db.query(main.ExposureHost).filter(main.ExposureHost.run_id == first.id).count() == 0
